from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from .cad_engine import build_manifest, load_step

MATERIAL_DENSITY_G_CM3 = {
    "6061": 2.70,
    "6061-T6": 2.70,
    "AL6061": 2.70,
    "304": 7.93,
    "SUS304": 7.93,
    "316": 7.98,
    "SUS316": 7.98,
    "Q235": 7.85,
    "A36": 7.85,
    "POM": 1.41,
}

DEFAULT_POLICY = {
    "currency": "CNY",
    "material": {"name": "UNKNOWN", "density_g_cm3": 0.0, "price_per_kg": 0.0},
    "process": {
        "stock_factor": 1.30,
        "machine_rate_per_hour": 120.0,
        "setup_minutes": 15.0,
        "removal_rate_cm3_min": 12.0,
        "inspection_minutes": 8.0,
        "tooling_pct": 0.08,
        "scrap_pct": 0.05,
    },
    "commercial": {"gross_margin_pct": 0.25},
    "assumptions": [
        "Material price is zero until explicitly supplied; missing material makes the quote incomplete.",
        "Default process rates are estimating assumptions, not market facts.",
        "All included parts are treated as CNC-machined unless excluded upstream.",
        "Cycle time is parametric, not CAM-simulated; production validation is required.",
    ],
}


def _merge(base: dict[str, Any], patch: dict[str, Any] | None) -> dict[str, Any]:
    out = json.loads(json.dumps(base))
    for key, value in (patch or {}).items():
        if isinstance(value, dict) and isinstance(out.get(key), dict):
            out[key] = _merge(out[key], value)
        elif value is not None:
            out[key] = value
    return out


def _number(value: Any, default: float, low: float, high: float) -> float:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return default
    return min(high, max(low, n))


def normalize_policy(policy: dict[str, Any] | None) -> dict[str, Any]:
    p = _merge(DEFAULT_POLICY, policy)
    m, pr, commercial = p["material"], p["process"], p["commercial"]
    m["name"] = str(m.get("name") or "UNKNOWN")[:80]
    density_default = next(
        (density for key, density in MATERIAL_DENSITY_G_CM3.items() if key.upper() in m["name"].upper()),
        0.0,
    )
    m["density_g_cm3"] = _number(m.get("density_g_cm3"), density_default, 0.0, 25)
    if m["density_g_cm3"] == 0 and density_default:
        m["density_g_cm3"] = density_default
    m["price_per_kg"] = _number(m.get("price_per_kg"), 0.0, 0, 100000)
    pr["stock_factor"] = _number(pr.get("stock_factor"), 1.3, 1.0, 20)
    pr["machine_rate_per_hour"] = _number(pr.get("machine_rate_per_hour"), 120, 0, 100000)
    pr["setup_minutes"] = _number(pr.get("setup_minutes"), 15, 0, 10000)
    pr["removal_rate_cm3_min"] = _number(pr.get("removal_rate_cm3_min"), 12, 0.01, 100000)
    pr["inspection_minutes"] = _number(pr.get("inspection_minutes"), 8, 0, 10000)
    pr["tooling_pct"] = _number(pr.get("tooling_pct"), 0.08, 0, 2)
    pr["scrap_pct"] = _number(pr.get("scrap_pct"), 0.05, 0, 2)
    commercial["gross_margin_pct"] = _number(commercial.get("gross_margin_pct"), 0.25, 0, 0.90)
    p["currency"] = str(p.get("currency") or "CNY")[:8].upper()
    p["assumptions"] = [str(x)[:300] for x in p.get("assumptions", [])][:50]
    return p


def calculate_costing(
    source_step: str | Path,
    project_id: str,
    part_ids: list[str],
    policy: dict[str, Any] | None = None,
) -> dict[str, Any]:
    p = normalize_policy(policy)
    manifest = build_manifest(load_step(source_step), project_id)
    selected_ids = set(part_ids)
    selected = [
        node for node in manifest["nodes"]
        if node["kind"] == "part" and node["id"] in selected_ids
    ]
    if not selected:
        raise ValueError("No valid parts selected for costing")

    groups: dict[str, list[dict[str, Any]]] = {}
    for part in selected:
        groups.setdefault(part["source_label_entry"], []).append(part)

    lines = []
    total_quote = total_cost = 0.0
    mat, proc, commercial = p["material"], p["process"], p["commercial"]

    for group in groups.values():
        part = group[0]
        qty = len(group)
        volume_mm3 = max(0.0, float(part.get("volume_mm3") or 0))
        mass_kg = volume_mm3 * mat["density_g_cm3"] / 1_000_000.0
        stock_mass_kg = mass_kg * proc["stock_factor"]
        material_cost = stock_mass_kg * mat["price_per_kg"]

        removed_cm3 = volume_mm3 * max(0.0, proc["stock_factor"] - 1.0) / 1000.0
        removal_minutes = removed_cm3 / proc["removal_rate_cm3_min"]
        cycle_minutes = proc["setup_minutes"] + removal_minutes + proc["inspection_minutes"]
        machine_cost = cycle_minutes / 60.0 * proc["machine_rate_per_hour"]
        tooling_cost = machine_cost * proc["tooling_pct"]

        base_cost = material_cost + machine_cost + tooling_cost
        risk_cost = base_cost * proc["scrap_pct"]
        unit_cost = base_cost + risk_cost
        margin = commercial["gross_margin_pct"]
        unit_quote = unit_cost / max(0.01, 1.0 - margin)
        extended_cost, extended_quote = unit_cost * qty, unit_quote * qty
        total_cost += extended_cost
        total_quote += extended_quote

        lines.append({
            "part_id": part["id"],
            "source_label_entry": part["source_label_entry"],
            "part_name": part["name"],
            "quantity": qty,
            "bbox_mm": part.get("bbox_mm", {}).get("size", [0, 0, 0]),
            "volume_mm3": volume_mm3,
            "mass_kg": mass_kg,
            "stock_mass_kg": stock_mass_kg,
            "material": mat["name"],
            "material_cost": material_cost,
            "setup_minutes": proc["setup_minutes"],
            "removal_minutes": removal_minutes,
            "inspection_minutes": proc["inspection_minutes"],
            "cycle_minutes": cycle_minutes,
            "machine_cost": machine_cost,
            "tooling_cost": tooling_cost,
            "risk_cost": risk_cost,
            "unit_cost": unit_cost,
            "unit_quote": unit_quote,
            "extended_quote": extended_quote,
            "confidence": "preliminary",
        })

    return {
        "schema_version": 1,
        "project_id": project_id,
        "currency": p["currency"],
        "policy": p,
        "line_count": len(lines),
        "occurrence_count": len(selected),
        "lines": lines,
        "totals": {"estimated_cost": total_cost, "quoted_price": total_quote},
        "quote_complete": bool(mat["density_g_cm3"] > 0 and mat["price_per_kg"] > 0),
        "release_status": "preliminary_requires_commercial_review",
    }


def render_bom_xlsx(costing: dict[str, Any], out_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM & Costing"
    headers = [
        "Part", "Qty", "Material", "BBox mm", "Mass kg", "Material Cost",
        "Cycle min", "Machine Cost", "Tooling", "Risk", "Unit Cost",
        "Unit Quote", "Extended Quote",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for line in costing["lines"]:
        ws.append([
            line["part_name"], line["quantity"], line["material"],
            " x ".join(f"{x:.2f}" for x in line["bbox_mm"]),
            line["mass_kg"], line["material_cost"], line["cycle_minutes"],
            line["machine_cost"], line["tooling_cost"], line["risk_cost"],
            line["unit_cost"], line["unit_quote"], line["extended_quote"],
        ])

    ws.append([])
    ws.append(["TOTAL QUOTE", costing["totals"]["quoted_price"], costing["currency"]])
    assumptions = wb.create_sheet("Assumptions")
    assumptions.append(["Key", "Value"])
    assumptions["A1"].font = assumptions["B1"].font = Font(bold=True)
    assumptions.append(["Policy", json.dumps(costing["policy"], ensure_ascii=False)])
    for i, item in enumerate(costing["policy"].get("assumptions", []), start=1):
        assumptions.append([f"Assumption {i}", item])
    wb.save(out_path)


def render_bom_csv(costing: dict[str, Any], out_path: str | Path) -> None:
    fields = [
        "part_name", "quantity", "material", "mass_kg", "material_cost",
        "cycle_minutes", "machine_cost", "unit_cost", "unit_quote", "extended_quote",
    ]
    with Path(out_path).open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for line in costing["lines"]:
            writer.writerow({key: line[key] for key in fields})


def render_quote_pdf(costing: dict[str, Any], out_path: str | Path) -> None:
    c = canvas.Canvas(str(out_path), pagesize=A4)
    width, height = A4
    y = height - 50
    c.setFont("Helvetica-Bold", 18)
    c.drawString(45, y, "PRELIMINARY MACHINING QUOTATION")
    y -= 24
    c.setFont("Helvetica", 9)
    c.drawString(45, y, f"Project: {costing['project_id']}")
    y -= 15
    c.drawString(45, y, f"Currency: {costing['currency']} | COMMERCIAL REVIEW REQUIRED")
    y -= 25

    c.setFont("Helvetica-Bold", 8)
    for x, title in [(45, "Part"), (260, "Qty"), (300, "Unit Quote"), (390, "Extended")]:
        c.drawString(x, y, title)
    y -= 12
    c.setFont("Helvetica", 8)

    for line in costing["lines"]:
        if y < 90:
            c.showPage()
            y = height - 50
            c.setFont("Helvetica", 8)
        c.drawString(45, y, line["part_name"][:34])
        c.drawRightString(285, y, str(line["quantity"]))
        c.drawRightString(375, y, f"{line['unit_quote']:.2f}")
        c.drawRightString(500, y, f"{line['extended_quote']:.2f}")
        y -= 12

    y -= 8
    c.setFont("Helvetica-Bold", 11)
    c.drawRightString(500, y, f"TOTAL: {costing['totals']['quoted_price']:.2f} {costing['currency']}")
    y -= 25
    c.setFont("Helvetica-Bold", 8)
    c.drawString(45, y, "Assumptions / limitations:")
    y -= 12
    c.setFont("Helvetica", 7)
    for item in costing["policy"].get("assumptions", []):
        c.drawString(55, y, ("- " + str(item))[:95])
        y -= 10
    c.save()


def generate_cost_bundle(
    source_step: str | Path,
    project_id: str,
    part_ids: list[str],
    out_dir: str | Path,
    policy: dict[str, Any] | None = None,
    revision: int = 0,
) -> dict[str, Any]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    costing = calculate_costing(source_step, project_id, part_ids, policy)
    costing["revision"] = revision
    paths = {
        "json": out_dir / "costing.json",
        "xlsx": out_dir / "bom.xlsx",
        "csv": out_dir / "bom.csv",
        "pdf": out_dir / "quotation.pdf",
    }
    paths["json"].write_text(json.dumps(costing, ensure_ascii=False, indent=2), encoding="utf-8")
    render_bom_xlsx(costing, paths["xlsx"])
    render_bom_csv(costing, paths["csv"])
    render_quote_pdf(costing, paths["pdf"])
    return {
        "summary": costing["totals"],
        "currency": costing["currency"],
        "quote_complete": costing["quote_complete"],
        "artifacts": {key: str(value) for key, value in paths.items()},
    }
